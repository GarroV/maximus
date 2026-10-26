"""Табель до эталона модуля 1 (T239): шапка сетки, подтверждение закрытия, загрузка.

Все три — экраны: проверяются живым запросом к странице, а не разбором шаблона.
Правила закрытия и разбор файла здесь не проверяются — их держат
`test_unit_closing.py` и тесты импорта.
"""
from __future__ import annotations

import io
import re

import pytest
from openpyxl import load_workbook

from conftest import body, login_as
from test_unit_closing import (  # noqa: F401 — фикстура clean_closures
    clean_closures,
    closable_unit,
    close_page,
    grid_url,
)

pytestmark = pytest.mark.usefixtures("clean_closures")


def test_grid_head_has_state_norm_file_and_unit_chips(client):
    login_as(client, "director")
    url = grid_url(client)
    html = body(client.get(url))

    assert 'class="status status--draft"' in html  # «Открыт»
    assert "sheet-head" in html
    assert f'href="{url}export/"' in html  # «Скачать xlsx»
    chips = re.findall(r'class="seg__item" href="[^"]*\?unit=([0-9a-f-]{36})"', html)
    assert len(chips) > 1, "у директора нет чипов точек"
    assert "<select" not in html.split("sheet-head", 1)[1].split("<table", 1)[0]


def test_unit_chip_narrows_the_grid_and_the_file(client):
    login_as(client, "director")
    url = grid_url(client)
    chips = re.findall(r'href="[^"]*\?unit=([0-9a-f-]{36})"', body(client.get(url)))
    unit = chips[0]

    full = body(client.get(url)).count('class="who"')
    one = body(client.get(f"{url}?unit={unit}")).count('class="who"')
    assert 0 < one < full

    book = load_workbook(io.BytesIO(client.get(f"{url}export/?unit={unit}").content))
    assert book.active.max_row - 1 == one, "в файле не то, что на экране"


def test_foreign_unit_in_the_address_is_ignored(client):
    login_as(client, "manager")
    url = grid_url(client)
    own = body(client.get(url)).count('class="who"')
    other = "00000000-0000-4000-8000-000000000000"
    assert body(client.get(f"{url}?unit={other}")).count('class="who"') == own


def test_export_is_an_xlsx_with_a_header_row(client):
    login_as(client, "accountant")
    response = client.get(f"{grid_url(client)}export/")
    assert response.status_code == 200
    assert "spreadsheetml" in response["Content-Type"]
    assert "attachment" in response["Content-Disposition"]
    assert load_workbook(io.BytesIO(response.content)).active.max_row > 1


def test_closing_asks_for_confirmation_before_closing(client):
    login_as(client, "manager")
    url = grid_url(client)
    listing = close_page(client, url)
    unit = closable_unit(listing)

    # Ссылка открывает подтверждение и сама ничего не закрывает.
    page = body(client.get(f"{url}close/?unit={unit}"))
    assert 'class="confirm"' in page and 'name="unit" value="' + unit in page
    assert "<script" not in page.split("<main", 1)[-1], "подтверждение не должно требовать JS"
    assert "is-closed" not in close_page(client, url), "подтверждение закрыло точку"


def test_confirmation_of_an_unknown_unit_falls_back_to_the_list(client):
    login_as(client, "manager")
    url = grid_url(client)
    page = body(client.get(f"{url}close/?unit=00000000-0000-4000-8000-000000000000"))
    assert 'class="confirm"' not in page and 'class="units"' in page


def test_upload_zone_is_a_native_file_input(client):
    login_as(client, "accountant")
    html = body(client.get(f"{grid_url(client)}import/"))
    assert re.search(r'<input class="dropzone" type="file" id="table" name="table"', html)
    assert '<label class="dropzone-title" for="table"' in html
