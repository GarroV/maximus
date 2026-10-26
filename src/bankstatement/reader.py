"""Файл банка → `Statement`. Формат угадывается по содержимому, а не по расширению.

Расширение врёт часто: «izvod.xml» из одного банка и из другого — разные
форматы, а xlsx бывает выгружен с расширением .xls. Поэтому XML узнаётся по
корню, таблица — по заголовкам, а книга Excel — по подписи zip-архива.

**Отказ — словами и сразу.** Файл, из которого не прочиталось ни одной
строки, не загружается «пустой выпиской»: пустая выписка выглядит как месяц без
движения денег, а это ложь, которую никто не заметит.
"""
from __future__ import annotations

import csv
import io
import unicodedata
import xml.parsers.expat
import zipfile
from functools import cache
from pathlib import Path
from xml.etree import ElementTree

import yaml
from django.utils.translation import gettext as _

from . import camt
from .model import IN, OUT, Line, Statement
from .values import amount, day

FORMATS = Path(__file__).with_name("formats.yaml")
# Выписка за месяц у самого оборотистого партнёра — сотни килобайт. Пять
# мегабайт — запас в десять раз, а не приглашение грузить архив за год.
MAX_BYTES = 5 * 1024 * 1024
XLSX_SIGNATURE = b"PK\x03\x04"
# Потолок строк — та же мысль, что у размера: месяц самого оборотистого
# партнёра — сотни операций, 20 000 — запас на годовую выгрузку, и не больше.
# Без потолка маленький файл мог бы развернуться в миллионы строк в памяти.
MAX_ROWS = 20_000
# Книга Excel — zip-архив. Лист выписки на сотни строк в развёрнутом виде —
# единицы мегабайт; 50 МБ — запас в десятки раз. Степень сжатия у живого
# листа — до нескольких десятков; выше 200 сжимаются только намеренно
# пустые данные, а не выписка.
# Ширина: в выписке банка — десяток-два колонок. Лист может объявить себя
# шириной в 16 384 колонки (XFD) или тянуть хвост пустых ячеек, поэтому дальше
# пятидесятой не читаем вовсе, а объявленному размеру листа не верим.
MAX_COLS = 50
# Общий бюджет непустых ячеек — последняя граница памяти, если и строки, и
# колонки в пределах: 20 000 строк по 10 заполненных полей с запасом.
MAX_CELLS = 400_000
MAX_UNPACKED_BYTES = 50 * 1024 * 1024
MAX_COMPRESSION_RATIO = 200
UTF16_BOMS = (b"\xff\xfe", b"\xfe\xff")
# Части книги, которые по имени — XML. Сравнение без учёта регистра.
XML_PART_SUFFIXES = (".xml", ".rels", ".vml")


class Unreadable(Exception):
    """Файл не разобран. Текст показывается человеку как есть."""


@cache
def formats() -> dict:
    return yaml.safe_load(FORMATS.read_text(encoding="utf-8"))


def read(content: bytes, file_name: str = "") -> Statement:
    if not content:
        raise Unreadable(_("Файл пустой."))
    if len(content) > MAX_BYTES:
        raise Unreadable(_("Файл больше 5 МБ — это не выписка за месяц."))
    if content.startswith(XLSX_SIGNATURE):
        statement = _table(_xlsx_rows(content))
    elif _looks_like_xml(content):
        statement = _xml(content)
    else:
        statement = _table(_csv_rows(content))
    if not statement.lines:
        raise Unreadable(_("В файле нет ни одной операции."))
    if len(statement.lines) > MAX_ROWS:
        raise Unreadable(
            _("В файле больше %(n)s операций — это не выписка за месяц.") % {"n": MAX_ROWS}
        )
    return statement


# --- XML ------------------------------------------------------------------------


def _looks_like_xml(content: bytes) -> bool:
    """XML в любой кодировке, которую понимает разборщик, — а не только в UTF-8."""
    if content.startswith(UTF16_BOMS):
        try:
            return content.decode("utf-16").lstrip().startswith("<")
        except UnicodeDecodeError:
            return False
    return content.removeprefix(b"\xef\xbb\xbf").lstrip()[:1] == b"<"


def _refuse_declarations(content: bytes, *, part: str = "") -> None:
    """Отказать XML с объявлением типа документа или сущностей.

    В выписке их не бывает, а в атаке на разборщик XML они есть всегда.
    Проверку делает сам разборщик (expat), а не поиск по байтам: он читает
    файл в той кодировке, которую файл объявил, и видит объявление там же,
    где его увидел бы `ElementTree`. Библиотека `defusedxml` делала бы то же
    одной строкой, но в зависимостях её нет.
    """
    def refuse(*_args):
        if part:
            raise Unreadable(_("В книге Excel часть %(part)s объявляет тип документа или "
                               "сущности — такая книга не принимается.") % {"part": part})
        raise Unreadable(_("XML с объявлением типа документа или сущностей не принимается."))

    probe = xml.parsers.expat.ParserCreate()
    probe.StartDoctypeDeclHandler = refuse
    probe.EntityDeclHandler = refuse
    probe.UnparsedEntityDeclHandler = refuse
    probe.ExternalEntityRefHandler = refuse
    try:
        probe.Parse(content, True)
    except xml.parsers.expat.ExpatError as error:
        if part:
            raise Unreadable(_("Книга Excel не открывается.")) from error
        raise Unreadable(_("XML не читается: %(why)s.") % {"why": error}) from error


def _xml(content: bytes) -> Statement:
    _refuse_declarations(content)
    try:
        root = ElementTree.fromstring(content)
    except ElementTree.ParseError as error:
        raise Unreadable(_("XML не читается: %(why)s.") % {"why": error}) from error
    if _local(root.tag) == "Document" and "camt.053" in root.tag:
        return camt.read(root)
    for profile in formats()["xml"]:
        found = root if _local(root.tag) == profile["root"] else _first(root, profile["root"])
        if found is not None:
            return _profiled(found, profile)
    raise Unreadable(_(
        "XML этого банка пока не знаком: знаем camt.053 и izvod iBank и Raiffeisen. "
        "Пришлите образец — формат заводится профилем."
    ))


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _first(node, name: str):
    for child in node.iter():
        if _local(child.tag) == name:
            return child
    return None


def _path(node, path: str) -> str:
    """Текст по пути «a/b» или атрибут «a@b» — без пространств имён."""
    element_path, _, attribute = path.partition("@")
    current = node
    for step in [s for s in element_path.split("/") if s]:
        current = next((c for c in current if _local(c.tag) == step), None)
        if current is None:
            return ""
    if attribute:
        return (current.get(attribute) or "").strip()
    return (current.text or "").strip()


def _profiled(root, profile: dict) -> Statement:
    fields = profile["fields"]
    by_attribute = profile.get("attributes", False)

    def value(record, name: str) -> str:
        if by_attribute:
            return (record.get(name) or "").strip()
        return _path(record, name)

    def first_of(record, names) -> str:
        for name in names:
            found = value(record, name)
            if found:
                return found
        return ""

    lines = []
    for number, record in enumerate(
        (r for r in root.iter() if _local(r.tag) == profile["record"]), start=1
    ):
        try:
            on = day(first_of(record, fields["date"]))
            if "amount" in fields:
                sum_ = amount(value(record, fields["amount"]))
                direction = _direction(value(record, fields["direction"]), profile)
            else:
                sum_, direction = _two_columns(
                    amount(value(record, fields["debit"])),
                    amount(value(record, fields["credit"])),
                )
        except ValueError as error:
            raise Unreadable(_("Операция %(n)s: %(why)s.") % {"n": number, "why": error}) from error
        if on is None or sum_ is None or direction is None:
            raise Unreadable(_("Операция %(n)s: нет даты, суммы или направления.") % {"n": number})
        def optional(name: str, record=record) -> str:
            return value(record, fields[name]) if fields.get(name) else ""

        lines.append(Line(
            op_date=on, direction=direction, amount=abs(sum_),
            party_name=optional("party_name"), party_account=optional("party_account"),
            purpose=optional("purpose"), bank_ref=optional("bank_ref"),
        ))
    account = _profile_account(root, profile.get("account", ""))
    currency = _profile_account(root, profile.get("currency", ""))
    return Statement(format=profile["code"], lines=lines, account=account, currency=currency,
                     period_from=min((x.op_date for x in lines), default=None),
                     period_to=max((x.op_date for x in lines), default=None))


def _profile_account(root, spec: str) -> str:
    """Значение шапки по профилю: «тег» или «тег@атрибут» — счёт, валюта."""
    if not spec:
        return ""
    tag, _, attribute = spec.partition("@")
    found = root if _local(root.tag) == tag else _first(root, tag)
    if found is None:
        return ""
    return (found.get(attribute) if attribute else found.text or "").strip()


def _direction(raw: str, profile: dict) -> str | None:
    values = profile.get("direction_values", {})
    text = raw.strip().casefold()
    if text in values.get("out", []):
        return OUT
    if text in values.get("in", []):
        return IN
    return None


def _two_columns(debit, credit):
    """«На терет» и «в корист» — две колонки; заполнена ровно одна."""
    debit = debit or 0
    credit = credit or 0
    if debit and not credit:
        return debit, OUT
    if credit and not debit:
        return credit, IN
    if debit and credit:
        raise ValueError(_("заполнены обе колонки — и расход, и приход"))
    return None, None


# --- таблица --------------------------------------------------------------------


def _csv_rows(content: bytes) -> list[list]:
    encodings = ("utf-16",) if content.startswith(UTF16_BOMS) else ()
    for encoding in (*encodings, "utf-8-sig", "cp1250", "cp1251"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise Unreadable(_("Кодировка файла не распознана: ждём UTF-8 или Windows-1250/1251."))
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
    except csv.Error:
        dialect = _Semicolon if sample.count(";") >= sample.count(",") else csv.excel
    try:
        return _capped(csv.reader(io.StringIO(text), dialect))
    except csv.Error as error:
        raise Unreadable(_("Таблица не читается: %(why)s.") % {"why": error}) from error


def _capped(rows) -> list[list]:
    """Строки с потолком по длине, ширине и числу ячеек — отказ словами.

    Ширина режется **до** сохранения строки, хвост пустых ячеек отбрасывается:
    иначе потолок строк обходился бы одной бесконечно широкой строкой.
    """
    taken = []
    cells = 0
    for row in rows:
        if len(taken) >= MAX_ROWS:
            raise Unreadable(_(
                "В файле больше %(n)s строк — это не выписка за месяц. "
                "Выгрузите выписку за более короткий период."
            ) % {"n": MAX_ROWS})
        kept = list(row[:MAX_COLS]) if isinstance(row, (list, tuple)) else list(row)[:MAX_COLS]
        while kept and (kept[-1] is None or kept[-1] == ""):
            kept.pop()
        cells += sum(1 for cell in kept if cell is not None and cell != "")
        if cells > MAX_CELLS:
            raise Unreadable(_("В файле больше %(n)s заполненных ячеек — "
                               "это не выписка за месяц.") % {"n": MAX_CELLS})
        taken.append(kept)
    return taken


class _Semicolon(csv.excel):
    delimiter = ";"


def _refuse_bombs(content: bytes) -> None:
    """Посмотреть в архив до распаковки: сколько он займёт развёрнутым."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            unpacked = sum(m.file_size for m in members)
            packed = sum(m.compress_size for m in members) or 1
            if unpacked > MAX_UNPACKED_BYTES or unpacked / packed > MAX_COMPRESSION_RATIO:
                raise Unreadable(_("Книга Excel разворачивается в слишком большой объём — "
                                   "это не выписка за месяц."))
            # Части книги — тот же XML, и правило у них то же, что у выписки
            # XML: объявлений типа документа и сущностей не бывает. Проверка —
            # после потолка объёма, поэтому читает не больше разрешённого.
            # Смотрим на содержимое, а не на имя: какую часть разборщик возьмёт
            # как XML, решает [Content_Types] книги, а не расширение, и регистр
            # имени ему тоже безразличен.
            for member in members:
                if member.is_dir():
                    continue
                body = archive.read(member)
                named_xml = member.filename.lower().endswith(XML_PART_SUFFIXES)
                if named_xml or _looks_like_xml(body):
                    _refuse_declarations(body, part=member.filename)
    except (zipfile.BadZipFile, OSError, RuntimeError) as error:
        raise Unreadable(_("Книга Excel не открывается.")) from error


def _xlsx_rows(content: bytes) -> list[list]:
    from openpyxl import load_workbook

    _refuse_bombs(content)
    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:  # openpyxl бросает что угодно на битом архиве
        raise Unreadable(_("Книга Excel не открывается.")) from error
    # Границы — свои, а не объявленные листом: размер листа (dimension) пишет
    # тот, кто сделал файл. Битый лист или лист-диаграмма openpyxl роняет чем
    # угодно и уже при обходе — отказ словами, а не падение страницы.
    try:
        sheet = book.worksheets[0]
        return _capped(sheet.iter_rows(max_row=MAX_ROWS + 1, max_col=MAX_COLS,
                                       values_only=True))
    except Unreadable:
        raise
    except Exception as error:  # noqa: BLE001 — openpyxl не обещает типа ошибки
        raise Unreadable(_("Лист книги Excel не читается: первым должен быть лист "
                           "с таблицей выписки.")) from error


def _header_key(text) -> str:
    plain = unicodedata.normalize("NFKD", str(text or "")).casefold()
    plain = "".join(ch for ch in plain if not unicodedata.combining(ch))
    return " ".join("".join(ch if ch.isalnum() else " " for ch in plain).split())


@cache
def _synonyms() -> dict[str, str]:
    known = {}
    for field, names in formats()["table"].items():
        for name in names:
            known[_header_key(name)] = field
    return known


def _find_header(rows: list[list]) -> tuple[int, dict[str, int], list[tuple[str, str]]]:
    """Строка заголовков: первая, где узнаны дата и хотя бы одна сумма.

    Ищем, а не берём первую: банки ставят над таблицей шапку — номер выписки,
    счёт, период, — и её строки заголовками не являются.
    """
    synonyms = _synonyms()
    for index, row in enumerate(rows[:30]):
        columns: dict[str, int] = {}
        mapping = []
        for position, cell in enumerate(row):
            field = synonyms.get(_header_key(cell))
            if field and field not in columns:
                columns[field] = position
                mapping.append((str(cell).strip(), field))
        has_sum = "amount" in columns or "debit" in columns or "credit" in columns
        if "date" in columns and has_sum:
            return index, columns, mapping
    raise Unreadable(_(
        "Не нашли колонки даты и суммы. Узнаём заголовки вроде «Datum valute», "
        "«Iznos na teret», «Iznos u korist» — проверьте, что в файле есть строка заголовков."
    ))


def _table(rows: list[list]) -> Statement:
    header, columns, mapping = _find_header(rows)

    def cell(row, field):
        position = columns.get(field)
        if position is None or position >= len(row):
            return None
        return row[position]

    lines = []
    for number, row in enumerate(rows[header + 1:], start=header + 2):
        if not any(str(c).strip() for c in row if c is not None):
            continue
        try:
            on = day(cell(row, "date"))
            if "amount" in columns:
                signed = amount(cell(row, "amount"))
                sum_, direction = (None, None) if signed is None else (
                    abs(signed), OUT if signed < 0 else IN)
            else:
                sum_, direction = _two_columns(amount(cell(row, "debit")),
                                               amount(cell(row, "credit")))
        except ValueError as error:
            raise Unreadable(_("Строка %(n)s: %(why)s.") % {"n": number, "why": error}) from error
        if on is None and sum_ is None:
            continue  # строка итога или подписи под таблицей
        if on is None or sum_ is None:
            raise Unreadable(_("Строка %(n)s: нет даты или суммы.") % {"n": number})

        def text(field, row=row):
            raw = cell(row, field)
            return "" if raw is None else str(raw).strip()

        lines.append(Line(
            op_date=on, direction=direction, amount=sum_,
            party_name=text("party_name"), party_account=text("party_account"),
            purpose=text("purpose"), bank_ref=text("bank_ref"),
        ))
    return Statement(format="table", lines=lines, mapping=mapping,
                     period_from=min((x.op_date for x in lines), default=None),
                     period_to=max((x.op_date for x in lines), default=None))
